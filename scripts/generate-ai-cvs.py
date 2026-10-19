"""Regenerate the two public AI CVs from one evidence-conscious content record."""

from pathlib import Path
from shutil import copyfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, HRFlowable


ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"
PORTFOLIO = "https://duytruyenthong.github.io/Alinguyen-portfolio/"
CONTACT = "Ho Chi Minh City, Vietnam  |  0902577827  |  nquangd1510@gmail.com"

PROFILE = (
    "Computer Science graduate building applied AI prototypes, customer-data workflows and "
    "digital products. Work spans retrieval and prompt design, a leakage-checked local churn "
    "baseline, public developer tooling and live commerce websites. I bring enterprise "
    "troubleshooting and client discovery to scoping, testing and explaining technical work."
)

SKILLS = [
    ("Applied AI", "LLM API and prompt workflows; retrieval-augmented generation prototypes; response review and documentation."),
    ("Data and evaluation", "Python, SQL, PostgreSQL, Pandas, scikit-learn, snapshot-date features and holdout evaluation."),
    ("Product engineering", "JavaScript, TypeScript, React, Next.js, Node.js, GitHub Actions, debugging and testing."),
    ("Business context", "Requirements discovery, commerce analytics, sales funnels and stakeholder communication."),
]

PROJECTS = [
    (
        "AI Knowledge Assistant | Retrieval and prompt prototype",
        [
            "Designed a retrieval-plus-generation flow for answering questions from internal knowledge sources.",
            "Prototyped prompt and retrieval steps and reviewed responses to representative knowledge questions. No production deployment or measured accuracy is claimed.",
        ],
    ),
    (
        "KKBox customer churn baseline | Local prototype",
        [
            "Built a snapshot-date feature pipeline and logistic-regression holdout baseline for customer churn exploration.",
            "Checked future-event exclusion with synthetic fixtures. Real KKBox model scores and business outcomes remain unmeasured.",
        ],
    ),
    (
        "Vibecode Product Kit | Public developer tool",
        [
            "Maintained an offline JavaScript CLI with Codex agent profiles, skills, templates and CI checks for reviewable product work.",
            "Public repository: github.com/DUYTRUYENTHONG/vibecode-product-kit. Adoption metrics are not claimed.",
        ],
    ),
    (
        "STOS Sales Traffic Operating System | Live product website",
        [
            "Built a public sales-page product with templates, a page studio, inquiry capture and an Insights workflow for views, actions and leads.",
            "AI suggestions are planned; customer adoption and business impact have not been independently measured.",
        ],
    ),
]

EXPERIENCE = [
    (
        "iTechWX | Microsoft Technical Support Specialist | Apr 2025 - Present",
        "Diagnose Microsoft 365 incidents across OneDrive, SharePoint, Exchange Online and Teams; reproduce symptoms, isolate causes, document evidence and coordinate escalations.",
    ),
    (
        "DUYTRUYENTHONG Agency | Founder and Operator | Feb 2024 - Present",
        "Translate client goals into websites, content systems and commerce workflows. Build and test internal product prototypes and use Codex-assisted development with reviewable QA records.",
    ),
    (
        "NINA Web Service | Solution Consultant | Dec 2023 - Dec 2024",
        "Gather requirements, define website architecture and workflows, and align client needs with delivery scope.",
    ),
    (
        "NINA Web Service | Web Developer | Jun 2023 - Dec 2023",
        "Developed React and JavaScript features; debugged and tested interfaces; prototyped a retrieval-based support chatbot.",
    ),
    (
        "Amazing Tech | Junior Front-End Developer | Sep 2022 - Jan 2023",
        "Supported responsive UI work, front-end refactoring, debugging and code-maintainability improvements.",
    ),
]


def build_pdf() -> None:
    output = PUBLIC / "cv-applied-ai-engineer.pdf"
    ink = colors.HexColor("#17212d")
    accent = colors.HexColor("#125d83")
    muted = colors.HexColor("#425466")
    styles = {
        "name": ParagraphStyle("Name", fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=ink, spaceAfter=4),
        "title": ParagraphStyle("Title", fontName="Helvetica-Bold", fontSize=9.5, leading=13, textColor=accent, spaceAfter=5),
        "contact": ParagraphStyle("Contact", fontName="Helvetica", fontSize=8, leading=11, textColor=muted, spaceAfter=3),
        "section": ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=9.5, leading=13, textColor=accent, spaceBefore=12, spaceAfter=5),
        "body": ParagraphStyle("Body", fontName="Helvetica", fontSize=8.6, leading=12.3, textColor=ink, spaceAfter=5),
        "item": ParagraphStyle("Item", fontName="Helvetica", fontSize=8.4, leading=11.8, textColor=ink, leftIndent=10, firstLineIndent=-8, spaceAfter=3),
        "heading": ParagraphStyle("Heading", fontName="Helvetica-Bold", fontSize=8.7, leading=12, textColor=ink, spaceBefore=5, spaceAfter=3),
    }
    story = []

    def para(text: str, style: str) -> None:
        story.append(Paragraph(escape(text), styles[style]))

    def section(title: str) -> None:
        para(title.upper(), "section")
        story.append(HRFlowable(width="100%", thickness=.5, color=colors.HexColor("#c7d9e2"), spaceAfter=6))

    para("NGUYEN QUANG DUY", "name")
    para("APPLIED AI ENGINEER | DATA WORKFLOWS | PRODUCT ENGINEERING", "title")
    para(CONTACT, "contact")
    para(f"Portfolio: {PORTFOLIO}", "contact")
    section("Profile")
    para(PROFILE, "body")
    section("Relevant skills")
    for label, detail in SKILLS:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(detail)}", styles["body"]))
    section("Selected AI and product work")
    for title, bullets in PROJECTS:
        para(title, "heading")
        for bullet in bullets:
            para(f"- {bullet}", "item")
    story.append(PageBreak())
    para("NGUYEN QUANG DUY", "name")
    section("Professional experience")
    for title, detail in EXPERIENCE:
        para(title, "heading")
        para(detail, "body")
    section("Education and languages")
    para("Bachelor of Computer Science | Vietnamese-German University | 2021 - 2026", "body")
    para("Selected coursework: Linear Algebra, Statistics and Discrete Mathematics.", "body")
    para("English: Professional Working Proficiency | German: A2", "body")
    section("Evidence and scope")
    para("Live websites and the public Vibecode repository can be inspected through the portfolio. The knowledge assistant and KKBox work are prototypes. No production AI deployment, KKBox model score or commercial uplift is claimed in this CV.", "body")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(17 * mm, 12 * mm, "NGUYEN QUANG DUY")
        canvas.drawRightString(A4[0] - 17 * mm, 12 * mm, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm, bottomMargin=19 * mm, title="Nguyen Quang Duy Applied AI Engineer CV", author="Nguyen Quang Duy")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_docx() -> None:
    output = PUBLIC / "cv-prompt-engineer.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(23, 33, 45)
    normal.paragraph_format.space_after = Pt(4)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Aptos Display"
        document.styles[style_name].font.color.rgb = RGBColor(18, 93, 131)
    document.styles["Title"].font.size = Pt(19)
    document.styles["Title"].font.color.rgb = RGBColor(23, 33, 45)
    document.styles["Heading 1"].font.size = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    document.styles["Heading 2"].font.size = Pt(9)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(6)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(2)

    document.add_paragraph("NGUYEN QUANG DUY", style="Title")
    document.add_paragraph("AI AND PROMPT ENGINEER | RAG PROTOTYPES | PYTHON")
    document.add_paragraph(CONTACT)
    document.add_paragraph(f"Portfolio: {PORTFOLIO}")
    document.add_heading("PROFILE", level=1)
    document.add_paragraph(PROFILE)
    document.add_heading("RELEVANT SKILLS", level=1)
    for label, detail in SKILLS:
        p = document.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(detail)
    document.add_heading("SELECTED AI AND PRODUCT WORK", level=1)
    for title, bullets in PROJECTS:
        document.add_heading(title, level=2)
        for bullet in bullets:
            p = document.add_paragraph(style="List Bullet")
            p.add_run(bullet)

    document.add_page_break()
    document.add_paragraph("NGUYEN QUANG DUY", style="Title")
    document.add_heading("PROFESSIONAL EXPERIENCE", level=1)
    for title, detail in EXPERIENCE:
        document.add_heading(title, level=2)
        document.add_paragraph(detail)
    document.add_heading("EDUCATION AND LANGUAGES", level=1)
    document.add_paragraph("Bachelor of Computer Science | Vietnamese-German University | 2021 - 2026")
    document.add_paragraph("Selected coursework: Linear Algebra, Statistics and Discrete Mathematics.")
    document.add_paragraph("English: Professional Working Proficiency | German: A2")
    document.add_heading("EVIDENCE AND SCOPE", level=1)
    document.add_paragraph("Live websites and the public Vibecode repository can be inspected through the portfolio. The knowledge assistant and KKBox work are prototypes. No production AI deployment, KKBox model score or commercial uplift is claimed in this CV.")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("NGUYEN QUANG DUY | AI AND PROMPT ENGINEERING")
    document.core_properties.title = "Nguyen Quang Duy AI and Prompt Engineer CV"
    document.core_properties.author = "Nguyen Quang Duy"
    document.save(output)


def build_business_pdfs() -> None:
    """Keep the two business download labels on one accurate, shared CV."""
    output = PUBLIC / "cv-business-product.pdf"
    accent = colors.HexColor("#125d83")
    ink = colors.HexColor("#17212d")
    styles = {
        "name": ParagraphStyle("BusinessName", fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=ink, spaceAfter=5),
        "title": ParagraphStyle("BusinessTitle", fontName="Helvetica-Bold", fontSize=9.2, leading=13, textColor=accent, spaceAfter=6),
        "body": ParagraphStyle("BusinessBody", fontName="Helvetica", fontSize=8.3, leading=11.5, textColor=ink, spaceAfter=4),
        "section": ParagraphStyle("BusinessSection", fontName="Helvetica-Bold", fontSize=9.3, leading=13, textColor=accent, spaceBefore=10, spaceAfter=4),
        "entry": ParagraphStyle("BusinessEntry", fontName="Helvetica-Bold", fontSize=8.5, leading=11.5, textColor=ink, spaceBefore=4, spaceAfter=2),
    }
    story = []

    def para(text: str, style: str = "body") -> None:
        story.append(Paragraph(escape(text), styles[style]))

    def section(title: str) -> None:
        para(title.upper(), "section")
        story.append(HRFlowable(width="100%", thickness=.5, color=colors.HexColor("#c7d9e2"), spaceAfter=5))

    para("NGUYEN QUANG DUY", "name")
    para("BUSINESS AND PRODUCT | DIGITAL COMMERCE | DATA-INFORMED DELIVERY", "title")
    para(CONTACT)
    para(f"Portfolio: {PORTFOLIO}")
    section("Profile")
    para("Founder and product operator with a Computer Science background and experience in digital marketing, solution consulting, enterprise support and applied AI prototypes. I turn client needs into clear requirements, commerce workflows and reviewable product builds. My data and AI work is presented with its measured and unmeasured limits.")
    section("Core competencies")
    para("Client discovery, business analysis, product scoping, sales funnels, website and content operations, SQL, Python, Power BI/DAX, stakeholder communication, testing and technical documentation.")
    section("Professional experience")
    for title, detail in EXPERIENCE:
        para(title, "entry")
        para(detail)
    section("Selected product and data work")
    para("STOS Sales Traffic Operating System | Live product website", "entry")
    para("Built a sales-page product with templates, a page studio, inquiry capture and Insights for views, actions and leads. AI suggestions remain planned; customer and revenue impact are not claimed.")
    para("Cho Kenh | Live marketplace website", "entry")
    para("Built a TikTok channel catalog with filters, risk and audit indicators, detail pages and manual Zalo inquiry paths. The intermediary process is manual; automated checkout and completed transactions are not claimed.")
    para("KKBox customer churn baseline | Local prototype", "entry")
    para("Explored retention features through a snapshot-date pipeline and logistic-regression holdout baseline. Synthetic fixtures check future-event exclusion; real KKBox model scores and revenue outcomes are unmeasured.")
    para("Vibecode Product Kit | Public developer tool", "entry")
    para("Maintained an offline CLI, Codex profiles, templates and CI checks for traceable planning and release reviews. Public repository: github.com/DUYTRUYENTHONG/vibecode-product-kit.")
    section("Education and languages")
    para("Bachelor of Computer Science | Vietnamese-German University | 2021 - 2026. English: Professional Working Proficiency; German: A2.")

    doc = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm, bottomMargin=16 * mm, title="Nguyen Quang Duy Business and Product CV", author="Nguyen Quang Duy")
    doc.build(story)
    copyfile(output, PUBLIC / "cv-business-development.pdf")


if __name__ == "__main__":
    build_pdf()
    build_docx()
    build_business_pdfs()
